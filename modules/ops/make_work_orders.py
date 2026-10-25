"""店铺调价/下架工单生成（跨平台，纯 openpyxl）。

从商品导出 xlsx 生成：
  - adjust-todo-<store>.csv   调价工单：已生效价格且现价>建议价(30%净利率)的 SKU，
                              按 倍数=现价/建议价 降序（grind_one/ego grind 的输入）
  - delist-todo-<store>.json  下架工单：亏损 SKU（现价<保本价）按 SKC 分组

定价模型（用户确认，见 HANDOFF-调价下架任务交接.md）：
  成本 = 单件成本 × 数量 + 9（运费+打包+设计/单）；广告抽申报价 20%
  保本价   = 成本 ÷ 0.8
  建议价30% = 成本 ÷ 0.5（向上取整到 0.5 元，与常用档 20件→24.5/100件→50/腕带24件→34 全部吻合）
  单件成本：餐巾类 0.16/张；身份腕带 0.33/条。其余类目成本未知 → 不生成（计入 skipped）

用法：python3 make_work_orders.py <store> [导出xlsx路径]
      store 如 store-c；缺省找 modules/collect/data/<store>-goods-*.xlsx 最新一份
⚠ 导出有滞后：执行前弹窗内现价/规格为准（grind 流程天然兼容）。
"""
import csv
import glob
import json
import math
import os
import re
import sys
from collections import defaultdict

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
BASE = os.path.dirname(os.path.abspath(__file__))
GOODS_DATA = os.path.abspath(os.path.join(BASE, "..", "collect", "data"))    # 商品导出（采集模块）
GRIND_DATA = os.path.abspath(os.path.join(BASE, "..", "grind", "data"))      # 调价工单输出
DELIST_DATA = os.path.abspath(os.path.join(BASE, "..", "delist", "data"))    # 下架工单输出

UNIT_COST = {"派对餐巾": 0.16, "餐巾": 0.16, "儿童派对餐巾": 0.16, "餐巾纸": 0.16,
             "身份腕带": 0.33}
FIXED_PER_ORDER = 9.0
QTY_RE = re.compile(r"(\d+)\s*(?:件|张|条|pcs|PCS|PC|pc|片)", re.I)


def parse_qty(row, idx):
    """数量：只认 规格1/2名称（「数量：50 件」）——它是 per-SKU 真值（与弹窗 SKU
    属性同源）。标题里的「50/100张」是整个链接的变体范围，按 SKU 取大者会错
    （实测：规格50件的 SKU 被标题读成100件 → 保本价虚高 → 假亏损）。
    两规格字段各解析出一个数量且不一致 → 视为歧义，返回 None 跳过。"""
    vals = []
    for k in ("规格1名称", "规格2名称"):
        v = row[idx[k]] if idx.get(k, -1) < len(row) else None
        if v:
            m = QTY_RE.search(str(v))
            if m:
                vals.append(int(m.group(1)))
    vals = list(dict.fromkeys(vals))
    if len(vals) == 1:
        return vals[0]
    if len(vals) > 1:
        return None  # 歧义
    return None  # 规格无数量：宁跳过不猜（标题数量是变体范围，不可靠）


def suggested(cost_total):
    """建议价30% = 成本合计÷0.5，再向上取整到 0.5 元。
    （20件→24.5 / 100件→50 / 腕带24条→34 / 101条→85 全部吻合）"""
    raw = cost_total / 0.5
    return math.ceil(raw / 0.5 - 1e-9) * 0.5


def load(fp):
    wb = openpyxl.load_workbook(fp, read_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    hdr = [str(h or "").strip() for h in next(it)]
    idx = {h: i for i, h in enumerate(hdr)}
    rows = [r for r in it if r is None or len(r) > idx.get("SKC ID", -1)]
    wb.close()
    return idx, rows


def cell(row, idx, key):
    i = idx.get(key, -1)
    return row[i] if 0 <= i < len(row) else None


def build(store, fp, dry=False, force=False):
    idx, rows = load(fp)
    adjust, loss_skus, skipped = [], [], defaultdict(int)
    seen = set()
    for r in rows:
        if r is None or cell(r, idx, "SKC ID") is None:
            continue
        skc = str(cell(r, idx, "SKC ID"))
        sku = str(cell(r, idx, "SKU ID"))
        if (skc, sku) in seen:
            continue
        seen.add((skc, sku))
        status = str(cell(r, idx, "商品状态") or "")
        pstat = str(cell(r, idx, "申报价格状态") or "")
        cat = str(cell(r, idx, "叶子类目名称") or "")
        if status != "在售中":
            skipped["非在售"] += 1
            continue
        if pstat != "已生效":
            skipped["价格非已生效"] += 1
            continue
        if cat not in UNIT_COST:
            skipped["类目成本未知"] += 1
            continue
        try:
            price = float(cell(r, idx, "申报价格(CNY)"))
        except (TypeError, ValueError):
            skipped["价格缺失"] += 1
            continue
        qty = parse_qty(r, idx)
        if not qty or qty <= 0:
            skipped["数量解析失败"] += 1
            continue
        cost_total = UNIT_COST[cat] * qty + FIXED_PER_ORDER
        break_even = round(cost_total / 0.8, 2)
        sug = suggested(cost_total)
        rec = {"skc": skc, "sku": sku, "price": price, "qty": qty, "cat": cat,
               "breakeven": break_even, "sug": sug, "title": str(cell(r, idx, "商品标题") or "")[:60]}
        if price < break_even:
            loss_skus.append(rec)          # 亏损 → 下架
        elif price > sug:
            rec["mult"] = round(price / sug, 2)
            adjust.append(rec)             # 高于建议价 → 调价
    adjust.sort(key=lambda x: -x["mult"])

    adj_fp = os.path.join(GRIND_DATA, f"adjust-todo-{store}.csv")
    del_fp = os.path.join(DELIST_DATA, f"delist-todo-{store}.json")
    if not dry:
        for p in (adj_fp, del_fp):
            if os.path.exists(p) and not force:
                raise SystemExit(f"⚠ {p} 已存在（可能是原版工单）。确认要覆盖请加 --force")
        with open(adj_fp, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(["SKC", "SKU", "现价", "建议价", "数量", "倍数"])
            for r in adjust:
                w.writerow([r["skc"], r["sku"], r["price"], r["sug"], r["qty"], r["mult"]])
        by_skc = defaultdict(list)
        for r in loss_skus:
            by_skc[r["skc"]].append(r)
        delist = [{"skc": int(k) if k.isdigit() else k,
                   "title": next(iter(v))["title"],
                   "skus": [[int(r["sku"]) if r["sku"].isdigit() else r["sku"],
                             r["price"], r["qty"]] for r in v],
                   "all": True}
                  for k, v in by_skc.items()]
        with open(del_fp, "w", encoding="utf-8") as f:
            json.dump(delist, f, ensure_ascii=False, indent=1)
    else:
        by_skc = defaultdict(list)
        for r in loss_skus:
            by_skc[r["skc"]].append(r)
        delist = [{"skc": k, "n": len(v)} for k, v in by_skc.items()]

    print(f"== {store}（{os.path.basename(fp)}）{' [dry-run]' if dry else ''} ==")
    print(f"  调价工单: {len(adjust)} SKU / {len(set(r['skc'] for r in adjust))} SKC"
          + ("" if dry else f" → {os.path.basename(adj_fp)}"))
    print(f"  下架工单: {len(delist)} SKC（{len(loss_skus)} SKU 亏损）"
          + ("" if dry else f" → {os.path.basename(del_fp)}"))
    print(f"  跳过: {dict(skipped)}")
    top = adjust[:8]
    if top:
        print("  倍数最高8行:")
        for r in top:
            print(f"    SKC={r['skc']} 数量={r['qty']} 现价={r['price']} → 建议 {r['sug']}（{r['mult']}x）")
    return adjust, delist


def main():
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    store = sys.argv[1]
    args = [a for a in sys.argv[2:] if a.startswith("--")]
    dry = "--dry" in args
    force = "--force" in args
    pos = [a for a in sys.argv[2:] if not a.startswith("--")]
    if pos:
        fp = pos[0]
    else:
        cands = sorted(glob.glob(os.path.join(GOODS_DATA, f"{store}-goods-*.xlsx")))
        if not cands:
            raise SystemExit(f"找不到 {store} 的导出 xlsx（modules/collect/data/{store}-goods-*.xlsx）")
        fp = cands[-1]
    build(store, fp, dry=dry, force=force)


if __name__ == "__main__":
    main()
